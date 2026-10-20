#!/usr/bin/env python3
"""
Normalise the Race #09 results sheet into the layout import_results.py reads.

The timing company delivered this race as a single "Results" sheet holding
four category blocks stacked on top of each other, each with its own header
row and a variable number of lap-split columns:

      Elite
      Bib | Name | Team Name | Lap 1 ... Lap 7 | Total time
      ...
      Trail Run
      Bib | Name | Team Name | Lap 1 ... Lap 4 | Total time

Every other race is one sheet per category with Place/Name/Team/Time columns.
Rather than teach the shared importer a second layout and risk the 85 events
that already import correctly, this rewrites the file into the familiar shape.
Finishing order is taken from row order, which the source is already sorted by;
riders with no total time are recorded as DNF with no rank.
"""
import os, re, sys
import openpyxl
from openpyxl import Workbook

SRC = sys.argv[1] if len(sys.argv) > 1 else (
    "public/2026 results/BBCh26Race09Results.xlsx")
DEST = sys.argv[2] if len(sys.argv) > 2 else (
    "/Users/parashramgore/Desktop/Dharma/Clients/BBch/archives/All Race Results"
    "/2026/Race #09 - XC/Result_ BBCh26 Race #09 - MTB XC & Trail Run.xlsx")

# Block heading in the source -> sheet name the importer already understands.
CATEGORY_SHEET = {
    "elite": "Elite",
    "xc open": "XC Open",
    "u09": "U-09",
    "trail run": "Trail Run",
}


def cells(row):
    return [("" if c is None else str(c).strip()) for c in row]


def main():
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [cells(r) for r in ws.iter_rows(values_only=True)]
    wb.close()

    blocks, current = [], None
    for r in rows:
        vals = [c for c in r if c]
        if not vals:
            continue
        low = [c.lower() for c in vals]

        # A lone heading cell starts a new category block.
        if len(vals) == 1 and low[0] in CATEGORY_SHEET:
            current = {"sheet": CATEGORY_SHEET[low[0]], "header": None, "rows": []}
            blocks.append(current)
            continue
        if current is None:
            continue
        if "bib" in low and "name" in low:
            current["header"] = r
            continue
        current["rows"].append(r)

    out = Workbook()
    out.remove(out.active)
    summary = []

    for b in blocks:
        if not b["header"] or not b["rows"]:
            continue
        hdr = [c.lower() for c in b["header"]]
        i_bib = hdr.index("bib")
        i_name = hdr.index("name")
        i_team = hdr.index("team name") if "team name" in hdr else None
        i_total = hdr.index("total time") if "total time" in hdr else len(hdr) - 1

        sheet = out.create_sheet(b["sheet"])
        sheet.append(["Place", "Bib", "Name", "Team Name", "Time"])

        place, finishers, dnf = 0, 0, 0
        for r in b["rows"]:
            def at(i):
                return r[i] if i is not None and i < len(r) else ""
            name = at(i_name)
            if not re.search(r"[A-Za-z]{2,}", name):
                continue
            total = at(i_total)
            # A missing or non-numeric total means they did not finish; the
            # sheet writes this as "DNF" or leaves the cell blank.
            finished = bool(re.match(r"^\d+:\d{2}:\d{2}", total))
            if finished:
                place += 1
                finishers += 1
                sheet.append([place, at(i_bib), name, at(i_team), total])
            else:
                dnf += 1
                sheet.append(["", at(i_bib), name, at(i_team), total or "DNF"])
        summary.append((b["sheet"], finishers, dnf))

    os.makedirs(os.path.dirname(DEST), exist_ok=True)
    out.save(DEST)

    print(f"  wrote {DEST}")
    total_f = total_d = 0
    for s, f, d in summary:
        print(f"    {s:<12} finishers={f:<3} dnf={d}")
        total_f += f
        total_d += d
    print(f"    {'TOTAL':<12} finishers={total_f:<3} dnf={total_d}  "
          f"rows={total_f + total_d}")


if __name__ == "__main__":
    main()
