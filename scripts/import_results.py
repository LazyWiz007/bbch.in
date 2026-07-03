import openpyxl, os, glob, re, json, datetime, unicodedata
from collections import Counter, defaultdict

BASE="/Users/parashramgore/Desktop/Dharma/Clients/BBch/archives/All Race Results"
OUT="/Users/parashramgore/Desktop/Dharma/Clients/BBch/website/src/data/generated"
os.makedirs(OUT, exist_ok=True)

# ---------- category mapping ----------
CAT_MAP = {}
def _m(canon, *variants):
    for v in variants: CAT_MAP[v.strip().lower()] = canon
_m("Elite","Elite","Open Men","Men Elite","ITT - Elite","ITT - ELITE","TTT - Elite","TTT - ELITE","Elite Cat Position")
_m("Amateur","Amateur","Amatuer","Amateur Open","Amateur Cat Position","ITT - Amateur","ITT - AMATEUR","TTT - Amateur","TTT - AMATEUR")
_m("Women","Women","Open Women","ITT - Women","ITT - WOMEN","TTT - Women","Women Cat Position")
_m("Men Master (40+)","Master","Masters","Men Master","Men Master (40+)","Open Master","Open Master (40+)","TTT - Master","TTT - MASTER","TTT - Men Master","TTT - Men Master (40+)","Master Cat Postion","Master Cat Position")
_m("Women Master (40+)","Women Master")
_m("Pioneer","Pioneer","Pioneer Open","Pioneer Category")
_m("Non-Road Bike","Non-Road Bike","Non Roadbike","Non Road Bike","Non-Road","TTT - Non-Road Bike")
_m("Open","XCO Open","XC Open","Open","XC - Open","All Categories","Amateur Open")  # note: Amateur Open handled above -> keep as Amateur; remove dup
CAT_MAP["amateur open"]="Amateur"
_m("U18","U-18","U18","Under-18","U-18 OPEN","U-18 Open","U18 OPEN","Open U-18","U18 Girls","MTB U-18")
_m("U16","U-16","U-16 OPEN","U16 Boys","XC - U16")
_m("U14","U-14","U14","U-14 Boys","U-14 Girls","U14 Girls")
_m("U12","U-12","U-12 Open")
_m("U11","U11 Open","U11 Boys","XC - U11")
_m("U10","U-10","U10","U10 - Boys","U10 - Girls")
_m("U09","U-09","U-09 Open")
_m("U-06 (FUN RIDE)","Kids","Kids Open")

# tabs we never import as results
EXCLUDE = set(x.lower() for x in [
 "Podium Winner","Podium","Podium Winners","PODIUM","Winners","Winner","Overall","Over All","OverAll","Overal",
 "All","Entire Results","Full Result","132Km","72 KM","100Km-Overall","60km-Overall","161 km- Overall","84km - Overall",
 "2017 PointsSystem","MTB TOP - 3","ROAD TOP - 3","MTB Elite Men","MTB Women","MTB Team","Road Team","Team Championship",
 "Trail Run - Men","Trail Run - Women","Trail Run Men","Climbs","Copy of Climbs","Updates",
 "Elite Team Position","Amateur Team Position","Master Team Position","Women Team Position",
])
GENERIC = set(x.lower() for x in ["Sheet1","Sheet2","Sheet3","Data"])
SKIP_DIR = re.compile(r"500km|Climbs Challenge|Season Points|Point System|CRIT SERIES", re.I)

# ---------- discipline detection ----------
def discipline_for(race_folder, tab_canon_raw):
    t = tab_canon_raw.lower()
    if t.startswith("ttt"): return "TTT"
    if t.startswith("itt"): return "ITT"
    if t.startswith("dh"): return "Downhill"
    if t.startswith("xc"): return None  # fall through to folder
    return None

def folder_discipline(name):
    n = name.lower()
    if "ttt" in n and "itt" in n: return "TTT/ITT"
    if "ttt" in n: return "TTT"
    if "itt" in n: return "ITT"
    if "crit" in n: return "Criterium"
    if re.search(r"\bdh\b|downhill|dh-xc|dh & ", n): return "Downhill/XC" if "xc" in n else "Downhill"
    if "cx" in n or "cyclocross" in n: return "Cyclocross"
    if any(k in n for k in ["classic","express","nandi epic","100k","road race","kolar","hassan"]): return "Road Race"
    if "nandi" in n and "itt" in n: return "ITT"
    if "dirt" in n: return "MTB (XC)"
    if "kiadb" in n or "circuit" in n: return "Criterium"
    if "xc" in n or "xco" in n or "mtb" in n: return "MTB (XC)"
    return "Road Race"

def series_for(name):
    n=name.lower()
    if "classic" in n: return "Bangalore Classic"
    if "express" in n: return "Express"
    if "nandi epic" in n: return "Nandi Epic"
    if "namma itt" in n or ("itt" in n and "ttt" not in n): return "Namma ITT"
    return None

# ---------- helpers ----------
def slugify(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii","ignore").decode()
    s = re.sub(r"[^a-zA-Z0-9]+","-", s).strip("-").lower()
    return s or "x"

def clean_name(raw):
    n = str(raw).strip()
    n = re.sub(r"\s+"," ", n)
    if n.isupper() or n.islower():
        n = n.title()
    return n

def norm_key(name):
    k = unicodedata.normalize("NFKD", name).encode("ascii","ignore").decode().lower()
    k = re.sub(r"[^a-z0-9 ]"," ", k)
    k = re.sub(r"\s+"," ", k).strip()
    return k

def parse_time(v):
    """returns (seconds_or_None, status, raw_str)"""
    if v is None: return None,"","" 
    if isinstance(v, datetime.time):
        return v.hour*3600+v.minute*60+v.second+v.microsecond/1e6,"OK",str(v)
    if isinstance(v, datetime.datetime):
        t=v.time(); return t.hour*3600+t.minute*60+t.second+t.microsecond/1e6,"OK",str(v)
    if isinstance(v,(int,float)):
        f=float(v)
        if 0<f<2: return f*86400,"OK",str(v)  # excel fractional day
        return None,"","" 
    s=str(v).strip()
    if s=="" : return None,"",""
    low=s.lower()
    if low in ("-","dnf","dns","dsq","dq","na","n/a"): return None, low.upper() if low!="-" else "DNF", s
    m=re.match(r"^(?:(\d+):)?(\d+):(\d+(?:\.\d+)?)$", s)
    if m:
        h=int(m.group(1) or 0); mi=int(m.group(2)); se=float(m.group(3))
        return h*3600+mi*60+se,"OK",s
    m=re.match(r"^(\d+):(\d+(?:\.\d+)?)$", s)  # MM:SS
    if m:
        return int(m.group(1))*60+float(m.group(2)),"OK",s
    return None,"?",s

def find_header(ws):
    for i,row in enumerate(ws.iter_rows(min_row=1,max_row=6,max_col=14,values_only=True)):
        vals=[str(c).strip().lower() if c is not None else "" for c in row]
        if "name" in vals:
            return i+1, vals
    return None,None

def col_idx(header, *names):
    for n in names:
        if n in header: return header.index(n)
    return None

def choose_file(rd):
    xlsx=glob.glob(os.path.join(rd,"*.xlsx"))
    res=[f for f in xlsx if os.path.basename(f).lower().startswith("result")]
    pool=res or xlsx
    def sc(f):
        b=os.path.basename(f).lower();s=0
        for t,w in (("(1)",3),("copy",4),("raw",5),("survey",6),("template",6),("responses",6)):
            if t in b:s+=w
        return s
    pool.sort(key=sc); return pool[0] if pool else None

# ---------- walk ----------
athletes={}   # key -> {names:Counter, teams:Counter, genders:Counter, id, slug}
events=[]
results=[]
issues=defaultdict(list)
def gender_for(cat, raw_tab):
    r=raw_tab.lower()
    if "women" in r or "girls" in r or cat in ("Women","Women Master (40+)"): return "F"
    return "M"

for year in sorted(os.listdir(BASE)):
    yd=os.path.join(BASE,year)
    if not os.path.isdir(yd) or not year.isdigit(): continue
    for race in sorted(os.listdir(yd)):
        rd=os.path.join(yd,race)
        if not os.path.isdir(rd): continue
        if SKIP_DIR.search(race):
            issues["skipped_dir"].append(f"{year}/{race}"); continue
        f=choose_file(rd)
        if not f:
            issues["no_xlsx"].append(f"{year}/{race}"); continue
        try:
            wb=openpyxl.load_workbook(f, read_only=True, data_only=True)
        except Exception as e:
            issues["open_error"].append(f"{year}/{race}: {e}"); continue

        m=re.match(r"(?:Race\s*#?\s*)?0*(\d+)", race)
        raceno=int(m.group(1)) if m else 99
        disc=folder_discipline(race)
        series=series_for(race)
        eid=f"{year}-{slugify(race)}"
        ev_cats=set(); ev_results=0

        # determine sheets to use
        sheets=wb.worksheets
        canon_sheets=[ws for ws in sheets if ws.title.strip().lower() in CAT_MAP]
        use=[]
        for ws in sheets:
            t=ws.title.strip(); tl=t.lower()
            if tl in EXCLUDE: continue
            if tl in CAT_MAP: use.append((ws, CAT_MAP[tl], t)); continue
            if tl.startswith("trail run"): continue
            if tl.startswith(("dh - ","seeding","final round")):
                use.append((ws, t, t)); continue  # keep DH format classes verbatim
            if (tl in GENERIC or "all categories" in tl) and not canon_sheets:
                use.append((ws, "Open", t)); continue
        for ws, canon, raw_tab in use:
            hr,header=find_header(ws)
            if not hr: continue
            ci_name=col_idx(header,"name")
            ci_place=col_idx(header,"place","position","pos","rank")
            ci_bib=col_idx(header,"bib","bib no","bib no.")
            ci_team=col_idx(header,"team name","team","team name ")
            ci_time=col_idx(header,"time","race time","finish time","overall time")
            if ci_name is None: continue
            disc2 = discipline_for(race, raw_tab) or disc
            for row in ws.iter_rows(min_row=hr+1, max_col=max(filter(None,[ci_name,ci_place,ci_bib,ci_team,ci_time]))+1, values_only=True):
                if ci_name>=len(row): continue
                nm=row[ci_name]
                if not nm or not isinstance(nm,str): 
                    if isinstance(nm,(int,float)): nm=str(nm)
                    else: continue
                nm=str(nm).strip()
                if not re.search(r"[A-Za-z]{2,}", nm): continue
                if nm.lower() in ("name","total","dnf","dns"): continue
                name=clean_name(nm)
                key=norm_key(name)
                if len(key)<2: continue
                team=row[ci_team] if ci_team is not None and ci_team<len(row) else None
                team=str(team).strip() if team and str(team).strip().upper()!="NONE" else None
                place=row[ci_place] if ci_place is not None and ci_place<len(row) else None
                bib=row[ci_bib] if ci_bib is not None and ci_bib<len(row) else None
                tval=row[ci_time] if ci_time is not None and ci_time<len(row) else None
                secs,status,raw=parse_time(tval)
                # rank
                rank=None; rstatus="OK"
                if isinstance(place,(int,float)): rank=int(place)
                elif isinstance(place,str) and place.strip().isdigit(): rank=int(place.strip())
                else:
                    rstatus = (str(place).strip().upper() if place and str(place).strip() not in ("-","") else (status if status in ("DNF","DNS","DSQ","DQ") else "DNF"))
                    if rstatus=="OK": rstatus="DNF"
                g=gender_for(canon, raw_tab)
                a=athletes.get(key)
                if not a:
                    a={"names":Counter(),"teams":Counter(),"genders":Counter()}
                    athletes[key]=a
                a["names"][name]+=1
                if team: a["teams"][team]+=1
                a["genders"][g]+=1
                try: bibn=int(bib) if bib is not None and str(bib).strip() not in ("","-") else None
                except: bibn=None
                results.append({
                    "eventId":eid,"athleteKey":key,"category":canon,"discipline":disc2,
                    "rank":rank,"status":None if rank else rstatus,
                    "timeSeconds":round(secs,2) if secs is not None else None,
                    "rawTime":raw or (str(tval) if tval is not None else ""),"bib":bibn,
                })
                ev_cats.add(canon); ev_results+=1
                if status=="?": issues["bad_time"].append(f"{year}/{race}/{raw_tab}: {raw}")
        wb.close()
        if ev_results==0:
            issues["empty_after_parse"].append(f"{year}/{race} ({[s.title for s in sheets]})"); 
        events.append({
            "id":eid,"year":int(year),"raceNo":raceno,"name":race,
            "discipline":disc,"series":series,"categories":sorted(ev_cats),
            "resultCount":ev_results,"resultsPublished":True,
        })

# finalize athletes
final_ath={}
used_slugs=Counter()
for key,a in athletes.items():
    disp=a["names"].most_common(1)[0][0]
    slug=slugify(disp); used_slugs[slug]+=1
    if used_slugs[slug]>1: slug=f"{slug}-{used_slugs[slug]}"
    gid=slug
    final_ath[key]={
        "id":gid,"slug":slug,"name":disp,
        "gender":a["genders"].most_common(1)[0][0],
        "team":a["teams"].most_common(1)[0][0] if a["teams"] else None,
        "aliases":[n for n,_ in a["names"].most_common() if n!=disp][:5],
    }
# map results athleteKey -> id
for r in results:
    r["athleteId"]=final_ath[r["athleteKey"]]["id"]
    del r["athleteKey"]

events=[e for e in events if e["resultCount"]>0]
good_ids=set(e["id"] for e in events)
results=[r for r in results if r["eventId"] in good_ids]
json.dump(list(final_ath.values()), open(os.path.join(OUT,"athletes.json"),"w"), ensure_ascii=False, indent=0)
json.dump(events, open(os.path.join(OUT,"events.json"),"w"), ensure_ascii=False, indent=0)
json.dump(results, open(os.path.join(OUT,"results.json"),"w"), ensure_ascii=False, indent=0)

# ---------- report ----------
print("=== IMPORT SUMMARY ===")
print("events:", len(events))
print("athletes (deduped):", len(final_ath))
print("results (rows):", len(results))
finishers=sum(1 for r in results if r["rank"])
print("  finishers:", finishers, " non-finishers(DNF/DNS):", len(results)-finishers)
print("results with time:", sum(1 for r in results if r["timeSeconds"] is not None))
by_disc=Counter(e["discipline"] for e in events)
print("\nevents by discipline:", dict(by_disc))
by_year=Counter(e["year"] for e in events)
print("events by year:", dict(sorted(by_year.items())))
print("\ncategories used:", dict(Counter(r["category"] for r in results)))
print("\ntop 12 athletes by #races:")
cnt=Counter(r["athleteId"] for r in results)
id2name={a["id"]:a["name"] for a in final_ath.values()}
for aid,c in cnt.most_common(12):
    print(f"   {c:3d}  {id2name[aid]}")
print("\n--- ISSUES ---")
for k,v in issues.items():
    print(f"{k}: {len(v)}")
    for x in v[:6]: print("    ",x)
